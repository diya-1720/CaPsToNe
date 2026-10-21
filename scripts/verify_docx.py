import os
import docx

docx_path = r"c:\Users\Anurag Singh\Desktop\capstone\CaPsToNe\AWEN_Project_Report_AF52.docx"
doc = docx.Document(docx_path)

print(f"=== AWEN REPORT AUDIT FOR {os.path.basename(docx_path)} ===")
print(f"Total Paragraphs: {len(doc.paragraphs)}")
print(f"Total Tables: {len(doc.tables)}")

total_words = 0
full_text = []
for p in doc.paragraphs:
    txt = p.text.strip()
    if txt:
        full_text.append(txt)
        total_words += len(txt.split())

for t in doc.tables:
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                txt = p.text.strip()
                if txt:
                    full_text.append(txt)
                    total_words += len(txt.split())

print(f"Total Word Count: {total_words:,} words")

# Check Students
students = [
    ("Diya Singh", "50"),
    ("Parth Topale", "58"),
    ("Nishtha Lokhande", "21"),
    ("Swara Patil", "38")
]
combined_text = "\n".join(full_text)
print("\n--- Student Verification ---")
for name, roll in students:
    found_name = name in combined_text
    found_roll = roll in combined_text
    print(f"{name} (Roll {roll}): Name Found = {found_name}, Roll Found = {found_roll}")

# Check Key Institutional Figures
print("\n--- Institutional Figures ---")
for inst in ["Dr. Kamal Shah", "Mrs. Joslyn Gracias", "St. John College of Engineering", "University of Mumbai", "2026-2027"]:
    print(f"'{inst}': {'FOUND' if inst.lower() in combined_text.lower() else 'MISSING'}")

# Check Chapters
print("\n--- Chapter Verification ---")
chapters = [
    "CHAPTER 1: INTRODUCTION",
    "CHAPTER 2: LITERATURE REVIEW",
    "CHAPTER 3: PROPOSED METHODOLOGY",
    "CHAPTER 4: PROJECT PLANNING",
    "CHAPTER 5: RESOURCES & FEASIBILITY",
    "CHAPTER 6: EXPECTED RESULTS",
    "CHAPTER 7: SUSTAINABILITY & IMPACT",
    "CHAPTER 8: CONCLUSION",
    "REFERENCES",
    "APPENDICES"
]
for ch in chapters:
    found = any(ch in p.text.upper() for p in doc.paragraphs)
    print(f"{ch}: {'FOUND' if found else 'MISSING'}")

# Check Image Shapes / Inline Shapes
print("\n--- Media & Figures Embedded ---")
import zipfile
with zipfile.ZipFile(docx_path) as z:
    media_files = [f for f in z.namelist() if f.startswith('word/media/')]
    print(f"Total embedded media files: {len(media_files)}")
    for mf in sorted(media_files):
        print(f"  {mf}: {z.getinfo(mf).file_size:,} bytes")

# Check for Forbidden Strings
print("\n--- Forbidden Token Audit ---")
for bad in ["lorem ipsum", "todo", "fixme", "100% accurate", "clinical diagnosis", "miraculous"]:
    count = combined_text.lower().count(bad)
    print(f"Token '{bad}': count = {count}")

print("\nAudit complete.")
